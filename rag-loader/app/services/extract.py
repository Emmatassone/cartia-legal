"""Extraccion de texto plano desde los formatos en que suelen venir los documentos."""

from __future__ import annotations

import logging
import re
import unicodedata
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}

# Ligaduras y comillas tipograficas que los PDF escaneados meten y que rompen
# tanto el matcheo de "ARTICULO" como el tsvector en espanol.
_REPLACEMENTS = {
    "\u00ad": "",  # guion suave
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u2013": "-",
    "\u2014": "-",
    "\ufb01": "fi",
    "\ufb02": "fl",
    "\u00a0": " ",
}

# Palabras que abren una unidad estructural de una norma argentina. Se usa como
# lookahead para no fusionar esas lineas con el parrafo anterior.
_STRUCTURAL_LOOKAHEAD = (
    r"\s*(?:ART[IÍ]CULO|ARTICULO|Art[íi]culo|ART\.|Art\.|"
    r"T[IÍ]TULO|T[íi]tulo|CAP[IÍ]TULO|Cap[íi]tulo|"
    r"SECCI[OÓ]N|Secci[óo]n|LIBRO|Libro|ANEXO|Anexo)"
)


def normalize_text(raw: str) -> str:
    text = unicodedata.normalize("NFKC", raw)
    for needle, replacement in _REPLACEMENTS.items():
        text = text.replace(needle, replacement)
    # Palabras cortadas por guion al final de linea (frecuente en PDF de boletines).
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Un salto de linea suelto dentro de un parrafo es ruido de layout, no estructura.
    # Se preservan los saltos que anteceden a un marcador estructural, porque el chunker
    # depende de que "ARTICULO 5" arranque su propia linea.
    text = re.sub(
        r"(?<![\n.:;])\n(?!" + _STRUCTURAL_LOOKAHEAD + r"|[\n\u2022\-\d])",
        " ",
        text,
    )
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            logger.warning("no se pudo extraer la pagina %s de %s: %s", index + 1, path.name, exc)
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def _extract_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text("\n")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = _extract_pdf(path)
    elif suffix == ".docx":
        raw = _extract_docx(path)
    elif suffix in {".html", ".htm"}:
        raw = _extract_html(path)
    elif suffix in {".txt", ".md"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"extension no soportada: {suffix} ({path.name})")

    text = normalize_text(raw)
    if len(text) < 50:
        raise ValueError(
            f"{path.name}: se extrajeron solo {len(text)} caracteres. "
            "Probablemente sea un PDF escaneado sin capa de texto (requiere OCR)."
        )
    return text
